from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.food import Dish, DishNutrient, PortionMeasure
from app.models.nutrient import Nutrient

INDB_URL = "https://www.anuvaad.org.in/wp-content/uploads/2020/07/Anuvaad_INDB_2024.11.xlsx"

# INDB column -> (nutrient code, source unit). INDB has no iodine or B12.
INDB_COLUMN_MAP: dict[str, tuple[str, str]] = {
    "energy_kcal": ("energy", "kcal"),
    "protein_g": ("protein", "g"),
    "fat_g": ("fat", "g"),
    "carb_g": ("carbohydrate", "g"),
    "freesugar_g": ("free_sugar", "g"),
    "fibre_g": ("fibre", "g"),
    "sfa_mg": ("saturated_fat", "mg"),
    "cholesterol_mg": ("cholesterol", "mg"),
    "calcium_mg": ("calcium", "mg"),
    "phosphorus_mg": ("phosphorus", "mg"),
    "magnesium_mg": ("magnesium", "mg"),
    "sodium_mg": ("sodium", "mg"),
    "potassium_mg": ("potassium", "mg"),
    "iron_mg": ("iron", "mg"),
    "copper_mg": ("copper", "mg"),
    "selenium_ug": ("selenium", "ug"),
    "manganese_mg": ("manganese", "mg"),
    "zinc_mg": ("zinc", "mg"),
    "vita_ug": ("vitamin_a", "ug"),
    "vite_mg": ("vitamin_e", "mg"),
    "vitk1_ug": ("vitamin_k", "ug"),
    "folate_ug": ("folate", "ug"),
    "vitb1_mg": ("thiamine", "mg"),
    "vitb2_mg": ("riboflavin", "mg"),
    "vitb3_mg": ("niacin", "mg"),
    "vitb5_mg": ("pantothenic_acid", "mg"),
    "vitb6_mg": ("vitamin_b6", "mg"),
    "vitb7_ug": ("biotin", "ug"),
    "vitc_mg": ("vitamin_c", "mg"),
}


def load_indb(db: Session, xlsx_path: Path) -> int:
    from etl.common import to_canonical

    nutrients = {n.code: n for n in db.scalars(select(Nutrient))}
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    rows = sheet.iter_rows(values_only=True)
    header = {name: index for index, name in enumerate(next(rows))}
    loaded = 0

    for row in rows:
        code = f"indb:{row[header['food_code']]}"
        dish = db.scalar(select(Dish).where(Dish.code == code))
        if dish is None:
            # Not editable: INDB ships no ingredient breakdown, so these are
            # reference entries and their intake counts as estimated.
            dish = Dish(code=code, source="INDB", is_editable=False)
            db.add(dish)
        dish.name = row[header["food_name"]]
        db.flush()

        for column, (nutrient_code, unit) in INDB_COLUMN_MAP.items():
            nutrient = nutrients.get(nutrient_code)
            value = row[header[column]] if column in header else None
            if nutrient is None or value in (None, ""):
                continue

            existing = db.scalar(
                select(DishNutrient).where(
                    DishNutrient.dish_id == dish.id,
                    DishNutrient.nutrient_id == nutrient.id,
                )
            )
            converted = to_canonical(float(value), unit, nutrient.unit)
            if existing is None:
                db.add(
                    DishNutrient(
                        dish_id=dish.id, nutrient_id=nutrient.id, per_100g=converted
                    )
                )
            else:
                existing.per_100g = converted

        serving_unit = row[header["servings_unit"]]
        serving_kcal = row[header["unit_serving_energy_kcal"]]
        per_100g_kcal = row[header["energy_kcal"]]
        # Guard against a zero/missing energy denominator (water, some
        # drinks) -- dividing by it would either crash or produce an
        # absurd gram weight, so a portion measure is simply skipped for
        # those rows rather than derived from a meaningless ratio.
        if serving_unit and serving_kcal and per_100g_kcal:
            grams = float(serving_kcal) * 100.0 / float(per_100g_kcal)
            if not db.scalar(
                select(PortionMeasure).where(
                    PortionMeasure.dish_id == dish.id,
                    PortionMeasure.measure_name == serving_unit,
                )
            ):
                db.add(
                    PortionMeasure(
                        dish_id=dish.id,
                        measure_name=str(serving_unit),
                        grams=grams,
                        is_default=True,
                    )
                )

        loaded += 1

    db.flush()
    return loaded
